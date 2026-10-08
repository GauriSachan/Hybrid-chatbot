"""
ingestion.py  --  Part 1: load -> clean -> chunk -> embed -> store

Public API (what other parts call):
    ingest_user_doc(path, session_id, original_name=None) -> str    # returns filename for Session.user_docs
    delete_session_docs(session_id)                                 # when a chat session ends
    build_kb(folder, manifest=None)                                 # one-off / rebuild KB

CLI:
    python ingestion.py build-kb ./kb_docs [--manifest manifest.json] [--reset]
    python ingestion.py stats
    python ingestion.py eval qa.json          # retrieval hit-rate check (see eval_retrieval)

manifest.json (optional) gives each KB file a title and URL for citations:
    {"Eiffel_Tower.txt": {"title": "Eiffel Tower", "url": "https://en.wikipedia.org/wiki/Eiffel_Tower"}}
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Optional

from kb_store import DATA_DIR, get_store

# Chunking parameters (Phase-I report: ~300-500 chars, ~50 overlap)
CHUNK_SIZE = 500
CHUNK_OVERLAP = 50
MIN_CHUNK_CHARS = 40
SUPPORTED = {".txt", ".md", ".pdf", ".docx"}

Page = tuple[Optional[int], str]      # (page_number or None, raw text)


# --------------------------------------------------------------------------- #
# 1. Loaders  ->  list of (page, text)
# --------------------------------------------------------------------------- #
def load_file(path: Path) -> list[Page]:
    ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        return [(None, path.read_text(encoding="utf-8", errors="ignore"))]
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        pages = [(i + 1, (p.extract_text() or "")) for i, p in enumerate(reader.pages)]
        return _strip_headers_footers(pages)
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:                       # keep table content (FAQ-style docs)
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return [(None, "\n\n".join(parts))]
    raise ValueError(f"Unsupported file type: {ext}")


def _strip_headers_footers(pages: list[Page]) -> list[Page]:
    """Drop lines repeated on >=50% of pages (running headers/footers) and bare page numbers."""
    norm = lambda l: re.sub(r"\d+", "#", l.strip().lower())
    n = len(pages)
    counts = Counter()
    for _, t in pages:
        counts.update({norm(l) for l in t.splitlines() if l.strip()})
    out = []
    for pg, t in pages:
        keep = []
        for l in t.splitlines():
            s = l.strip()
            if re.fullmatch(r"(page\s*)?\d+(\s*(of|/)\s*\d+)?", s, re.I):
                continue
            if n >= 3 and len(s) < 100 and counts[norm(l)] >= max(3, n * 0.5):
                continue
            keep.append(l)
        out.append((pg, "\n".join(keep)))
    return out


# --------------------------------------------------------------------------- #
# 2. Cleaning
# --------------------------------------------------------------------------- #
def clean_text(t: str) -> str:
    t = unicodedata.normalize("NFKC", t).replace("\x00", " ")
    t = re.sub(r"-\n(?=[a-z])", "", t)               # de-hyphenate line breaks
    t = re.sub(r"[ \t\f\v]+", " ", t)
    t = re.sub(r"\n\s*\n+", "\n\n", t)               # paragraph breaks
    t = re.sub(r"(?<!\n)\n(?!\n)", " ", t)           # soft line breaks -> space
    t = re.sub(r"\[\d+\]", "", t)                    # wiki-style [12] reference marks
    return re.sub(r" {2,}", " ", t).strip()


# --------------------------------------------------------------------------- #
# 3. Chunking (sentence-aware, with character overlap)
# --------------------------------------------------------------------------- #
_SENT = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9\"'(\[])")


def chunk_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP,
               min_len: int = MIN_CHUNK_CHARS) -> list[str]:
    sents: list[str] = []
    for para in text.split("\n\n"):
        for s in _SENT.split(para):
            s = s.strip()
            while len(s) > size:                      # very long sentence: split on whitespace
                cut = s.rfind(" ", 0, size)
                cut = cut if cut > size * 0.5 else size
                sents.append(s[:cut].strip())
                s = s[cut:].strip()
            if s:
                sents.append(s)

    chunks: list[str] = []
    cur: list[str] = []
    cur_len, fresh = 0, 0                             # fresh = new sentences (not overlap)
    for s in sents:
        if cur and cur_len + len(s) + 1 > size:
            chunks.append(" ".join(cur))
            tail = chunks[-1][-overlap:] if overlap else ""
            tail = tail.split(" ", 1)[1] if " " in tail else ""   # start tail at a word boundary
            cur, cur_len, fresh = ([tail] if tail else []), len(tail), 0
        cur.append(s)
        cur_len += len(s) + 1
        fresh += 1
    if cur and fresh:
        chunks.append(" ".join(cur))
    chunks = [c for c in chunks if len(c) >= min_len]
    return chunks


def make_chunks(pages: list[Page]) -> list[dict]:
    """Chunk page by page so every chunk keeps its own page number."""
    out = []
    for pg, raw in pages:
        for c in chunk_text(clean_text(raw)):
            out.append({"text": c, "page": pg})
    return out


# --------------------------------------------------------------------------- #
# 4. Ingest
# --------------------------------------------------------------------------- #
def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def _slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", s)[:80]


def ingest_file(path: str | Path, origin: str, session_id: Optional[str] = None,
                source: Optional[str] = None, url: Optional[str] = None) -> dict:
    path = Path(path)
    if path.suffix.lower() not in SUPPORTED:
        raise ValueError(f"Unsupported file type '{path.suffix}'. Supported: {sorted(SUPPORTED)}")
    store = get_store()
    sha = _sha256(path)
    dup = store.has_document(sha, origin, session_id)
    if dup:
        return {"doc_id": dup, "source": source or path.name, "chunks": 0, "status": "duplicate_skipped"}

    chunks = make_chunks(load_file(path))
    if not chunks:
        return {"doc_id": None, "source": source or path.name, "chunks": 0,
                "status": "no_text_extracted"}       # e.g. scanned PDF (needs OCR)
    doc_id = f"{origin}-{_slug(session_id or 'kb')}-{_slug(path.stem)}-{sha[:8]}"
    n = store.add_document(doc_id, origin, source or path.name, chunks, url=url,
                           session_id=session_id, sha256=sha)
    return {"doc_id": doc_id, "source": source or path.name, "chunks": n, "status": "ok"}


def ingest_user_doc(path: str | Path, session_id: str, original_name: Optional[str] = None) -> str:
    """Call from the /upload endpoint. Returns the filename to append to Session.user_docs.
    Raises ValueError if the file is unsupported or has no extractable text."""
    name = original_name or Path(path).name
    r = ingest_file(path, "user_docs", session_id=session_id, source=name)
    if r["status"] == "no_text_extracted":
        raise ValueError(f"No readable text found in '{name}' (scanned/image-only PDF?).")
    return name


def delete_session_docs(session_id: str) -> int:
    return get_store().delete_session(session_id)


def build_kb(folder: str | Path, manifest: Optional[str | Path] = None) -> list[dict]:
    folder = Path(folder)
    meta = json.loads(Path(manifest).read_text(encoding="utf-8")) if manifest else {}
    results = []
    for p in sorted(folder.rglob("*")):
        if p.is_file() and p.suffix.lower() in SUPPORTED:
            m = meta.get(p.name, {})
            r = ingest_file(p, "kb", source=m.get("title") or p.stem.replace("_", " "), url=m.get("url"))
            print(f"[{r['status']:<18}] {p.name}  ({r['chunks']} chunks)")
            results.append(r)
    print("KB stats:", get_store().stats())
    return results


# --------------------------------------------------------------------------- #
# 5. Validation helper (hand the same qa.json to Part 5)
# --------------------------------------------------------------------------- #
def eval_retrieval(qa_path: str, top_k: int = 5) -> float:
    """qa.json: [{"question": "...", "expect": "substring that must appear in a retrieved chunk"}]
    Add {"question": "...", "expect": null} for out-of-scope questions: reports their top score."""
    from hybrid_retriever import HybridRetriever
    r = HybridRetriever(get_store())
    qa = json.loads(Path(qa_path).read_text(encoding="utf-8"))
    hits = total = 0
    for item in qa:
        res = r.search(item["question"], top_k=top_k, sources=["kb"])
        if item.get("expect") is None:
            print(f"[OOS ] top score={res[0].score:.2f}  {item['question']}" if res else f"[OOS ] none  {item['question']}")
            continue
        total += 1
        ok = any(item["expect"].lower() in c.text.lower() for c in res)
        hits += ok
        print(f"[{'HIT ' if ok else 'MISS'}] {item['question']}")
    rate = hits / total if total else 0.0
    print(f"hit@{top_k}: {hits}/{total} = {rate:.2%}")
    return rate


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build-kb"); b.add_argument("folder"); b.add_argument("--manifest"); b.add_argument("--reset", action="store_true")
    sub.add_parser("stats")
    e = sub.add_parser("eval"); e.add_argument("qa"); e.add_argument("--k", type=int, default=5)
    a = ap.parse_args()
    if a.cmd == "build-kb":
        if a.reset:
            shutil.rmtree(DATA_DIR, ignore_errors=True)
        build_kb(a.folder, a.manifest)
    elif a.cmd == "stats":
        print(get_store().stats())
    elif a.cmd == "eval":
        eval_retrieval(a.qa, a.k)
